# -*- coding: utf-8 -*-
"""Merge character/skill/carta tables + download portraits into one summary."""
from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PAGES = os.path.join(ROOT, "docs", "reference", "gamekee", "pages")
TABLES = os.path.join(ROOT, "docs", "reference", "gamekee", "tables")
USER_OUT = os.path.join(ROOT, "天命之子数据")

IMG_TAG = re.compile(r"<img\b([^>]*)>", re.I)
ATTR = re.compile(r"""([\w:-]+)\s*=\s*(['"])(.*?)\2""", re.S)
URL_W = re.compile(r"/w_(\d+)/")
URL_H = re.compile(r"/h_(\d+)/")

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
HEADERS = {
    "User-Agent": UA,
    "Referer": "https://www.gamekee.com/dc/",
    "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
}

CHAR_FIELDS = [
    "content_id", "avatar_file", "avatar_url", "name", "rarity", "element", "role",
    "cp_init", "hp_init", "atk_init", "def_init", "agl_init", "crt_init",
    "cp_max", "hp_max", "atk_max", "def_max", "agl_max", "crt_max",
    "auto_text", "tap_text", "tap_ignited", "slide_text", "slide_ignited",
    "drive_text", "drive_ignited", "leader_text", "wiki_url",
]


def abs_url(u: str) -> str:
    u = (u or "").strip()
    if not u:
        return ""
    if u.startswith("//"):
        return "https:" + u
    if u.startswith("http://"):
        return "https://" + u[len("http://"):]
    if u.startswith("https://"):
        return u
    return "https://" + u.lstrip("/")


def url_wh(url: str):
    mw = URL_W.search(url or "")
    mh = URL_H.search(url or "")
    return (int(mw.group(1)) if mw else 0, int(mh.group(1)) if mh else 0)


def parse_imgs(content: str):
    out = []
    for m in IMG_TAG.finditer(content or ""):
        attrs = {k.lower(): v for k, _, v in ATTR.findall(m.group(1))}
        src = attrs.get("data-real") or attrs.get("src") or attrs.get("data-src") or ""
        if not src:
            continue
        try:
            w = int(attrs.get("data-width") or 0)
        except ValueError:
            w = 0
        try:
            h = int(attrs.get("data-height") or 0)
        except ValueError:
            h = 0
        out.append((src, w, h))
    return out


def score_head(url: str, w: int, h: int) -> float:
    u = (url or "").strip()
    if not u:
        return -1
    low = u.lower()
    if ".gif" in low.split("?")[0]:
        return -1
    uw, uh = url_wh(u)
    w = w or uw
    h = h or uh
    if "w_100/h_25" in low or (w == 100 and h == 25) or (uw == 100 and uh == 25):
        return -1
    if h and h <= 40 and w and w >= 80:
        return -1
    # CDN resize query used as wiki card icon
    if "h_250,w_250" in low or "h_250%2cw_250" in low:
        return 96
    if "/w_250/h_250/" in low:
        return 100
    if "988768.png" in low:
        return -1
    if "/w_135/h_135/" in low:
        return 96
    if "/w_128/h_128/" in low:
        return 35
    if any(x in low for x in ("/w_246/", "/w_247/", "/w_248/", "/w_249/")):
        return 94
    if w and h:
        ratio = w / float(h)
        square = 0.82 <= ratio <= 1.22
        if square and 110 <= w <= 280:
            return 92 + min(w, 250) / 250.0
        if square and w == 512:
            return 72
        if square and 80 <= w < 110:
            return 55
        if w >= 400 or h >= 400:
            return 8
    if low.split("?")[0].endswith(".png"):
        return 45
    return 12


def pick_from_content(content: str):
    best = ("", -1.0)
    for src, w, h in parse_imgs(content):
        s = score_head(src, w, h)
        if s > best[1]:
            best = (src, s)
        if s >= 97:
            break
    if best[1] >= 40:
        return abs_url(best[0]), best[1]
    return "", -1


def load_json(path: str):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def collect_related_icons():
    by_id = {}
    by_name = {}
    if not os.path.isdir(PAGES):
        return by_id, by_name
    for fn in os.listdir(PAGES):
        if not fn.endswith(".json"):
            continue
        try:
            data = load_json(os.path.join(PAGES, fn)).get("data") or {}
        except Exception:
            continue
        for rs in data.get("related_strategies") or []:
            if not isinstance(rs, dict):
                continue
            icon = (rs.get("icon") or "").strip()
            if not icon:
                continue
            iid = rs.get("id")
            name = (rs.get("name") or "").strip()
            if iid:
                by_id[int(iid)] = abs_url(icon)
            if name:
                by_name[name] = abs_url(icon)
    return by_id, by_name


def pick_avatar(cid: int, name: str, related_id, related_name):
    path = os.path.join(PAGES, "%d.json" % cid)
    content = ""
    if os.path.isfile(path):
        try:
            data = load_json(path).get("data") or {}
            content = data.get("content") or ""
            models = data.get("model_list") or []
            if models and not content:
                content = (models[0] or {}).get("html") or ""
        except Exception:
            content = ""
    url, score = pick_from_content(content)
    if score >= 90:
        return url, "content"
    rel = related_id.get(cid) or related_name.get(name) or ""
    if rel and score_head(rel, 0, 0) >= 40:
        if not url or score_head(rel, 0, 0) >= score:
            return rel, "related"
    if url:
        return url, "content-weak"
    return "", "missing"


def read_csv(path: str):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def ext_of(url: str) -> str:
    base = (url or "").split("?")[0].lower()
    for e in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        if base.endswith(e):
            return e
    return ".png"


def is_image_bytes(data: bytes) -> bool:
    if len(data) < 12:
        return False
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return True
    if data[:2] == b"\xff\xd8":
        return True
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return True
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return True
    return False


def url_alts(url: str):
    seen = set()
    out = []
    cands = [url]
    cands.append(url.replace("cdnimg.gamekee.com", "cdnimg-v2.gamekee.com"))
    cands.append(url.replace("cdnimg-v2.gamekee.com", "cdnimg.gamekee.com"))
    cands.append(url.replace("cdnimg01.gamekee.com", "cdnimg-v2.gamekee.com"))
    cands.append(url.replace("cdnimg01.gamekee.com", "cdnimg.gamekee.com"))
    if "?" in url:
        bare = url.split("?")[0]
        cands.append(bare)
        cands.append(bare.replace("cdnimg.gamekee.com", "cdnimg-v2.gamekee.com"))
        cands.append(bare.replace("cdnimg-v2.gamekee.com", "cdnimg.gamekee.com"))
    for u in cands:
        if u and u not in seen:
            seen.add(u)
            out.append(u)
    return out


def fetch_image(url: str) -> bytes:
    req = Request(url, headers=HEADERS)
    with urlopen(req, timeout=25) as resp:
        data = resp.read()
    if not is_image_bytes(data):
        raise RuntimeError("not-image")
    return data


def download_one(url: str, dest: str) -> str:
    if os.path.isfile(dest) and os.path.getsize(dest) > 64:
        return "exists"
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    last = "fail"
    for u in url_alts(url):
        try:
            data = fetch_image(u)
            tmp = dest + ".part"
            with open(tmp, "wb") as f:
                f.write(data)
            os.replace(tmp, dest)
            return "ok"
        except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as e:
            last = "fail:%s" % e
            continue
    return last


def download_all(items, avatars_dir):
    """items: list of dict with content_id + avatar_url. Mutates avatar_file."""
    jobs = []
    for rec in items:
        url = rec.get("avatar_url") or ""
        cid = rec["content_id"]
        if not url:
            rec["avatar_file"] = ""
            continue
        fn = "%s%s" % (cid, ext_of(url))
        rec["avatar_file"] = "avatars/" + fn
        dest = os.path.join(avatars_dir, fn)
        jobs.append((url, dest, rec))

    stats = Counter()
    if not jobs:
        return stats
    n = len(jobs)
    done = 0
    t0 = time.time()
    with ThreadPoolExecutor(max_workers=10) as ex:
        futs = {ex.submit(download_one, url, dest): (url, dest, rec) for url, dest, rec in jobs}
        for fut in as_completed(futs):
            url, dest, rec = futs[fut]
            try:
                st = fut.result()
            except Exception as e:
                st = "fail:%s" % e
            if st.startswith("fail"):
                rec["avatar_file"] = ""
            stats[st.split(":")[0]] += 1
            done += 1
            if done % 40 == 0 or done == n:
                print("  下载 %d/%d  %.1fs  %s" % (done, n, time.time() - t0, dict(stats)), flush=True)
    return stats


def write_csv(path, rows, fields):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})


def e(s) -> str:
    return html.escape("" if s is None else str(s), quote=True)


def skill_cell(text, ignited=""):
    parts = []
    if text:
        parts.append('<div class="sk">%s</div>' % e(text))
    if ignited:
        parts.append('<div class="sk ign">点火U：%s</div>' % e(ignited))
    return "".join(parts) or '<span class="empty">—</span>'


def stat_block(row, suffix):
    keys = [
        ("CP", "cp_" + suffix),
        ("HP", "hp_" + suffix),
        ("攻", "atk_" + suffix),
        ("防", "def_" + suffix),
        ("敏", "agl_" + suffix),
        ("暴", "crt_" + suffix),
    ]
    bits = []
    any_v = False
    for lab, k in keys:
        v = (row.get(k) or "").strip()
        if v:
            any_v = True
        bits.append("<span><i>%s</i>%s</span>" % (lab, e(v or "—")))
    return '<div class="stats %s">%s</div>' % ("ok" if any_v else "na", "".join(bits))


def el_class(el):
    return {
        "火": "el-fire", "水": "el-water", "木": "el-wood",
        "光": "el-light", "暗": "el-dark",
    }.get(el, "el-none")


def avatar_html(row, size="md"):
    local = row.get("avatar_file") or ""
    remote = row.get("avatar_url") or ""
    src = local if local else remote
    if not src:
        return '<div class="ph">无</div>'
    fallback = remote if (local and remote and remote != local) else ""
    return (
        '<img class="av %s" src="%s" alt="%s" loading="lazy" data-r="%s" '
        'onerror="if(!this.dataset.fb&&this.dataset.r){this.dataset.fb=1;this.src=this.dataset.r;}else{this.replaceWith(ph());}">'
        % (size, e(src), e(row.get("name") or ""), e(fallback))
    )


CSS = r"""
:root {
  --bg:#10131a; --panel:#181c27; --line:#2a3144; --text:#e8ecf4; --muted:#8b93a7;
  --fire:#e85d4c; --water:#4aa3e8; --wood:#5cbf6a; --light:#e8c84a; --dark:#9b7ae8;
}
* { box-sizing:border-box; }
html,body { margin:0; background:var(--bg); color:var(--text);
  font:14px/1.45 "Segoe UI","Microsoft YaHei",sans-serif; }
a { color:#8ec4ff; }
header {
  position:sticky; top:0; z-index:20;
  background:linear-gradient(180deg,#1a2030 0%,#141824 100%);
  border-bottom:1px solid var(--line); padding:12px 18px 10px;
}
h1 { margin:0 0 8px; font-size:20px; font-weight:700; }
.sub { color:var(--muted); font-size:12px; margin-bottom:10px; }
.tabs { display:flex; gap:8px; flex-wrap:wrap; margin-bottom:8px; }
.tabs button, .views button {
  background:#232a3b; color:var(--text); border:1px solid var(--line);
  border-radius:8px; padding:6px 12px; cursor:pointer;
}
.tabs button.on, .views button.on { background:#3a4a78; border-color:#6a82c4; }
.filters { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
.filters input, .filters select {
  background:#0e121c; color:var(--text); border:1px solid var(--line);
  border-radius:6px; padding:6px 8px;
}
.filters input { min-width:220px; }
.meta { color:var(--muted); font-size:12px; margin-left:auto; }
main { padding:12px 16px 40px; }
.panel { display:none; }
.panel.on { display:block; }
.table-wrap { overflow:auto; max-height:calc(100vh - 170px);
  border:1px solid var(--line); border-radius:10px; }
table { border-collapse:collapse; width:100%; min-width:1400px; }
th, td { border-bottom:1px solid var(--line); padding:8px 8px; vertical-align:top; }
th { position:sticky; top:0; background:#1e2433; text-align:left; font-size:12px;
  color:#c5cce0; z-index:5; }
tr:hover td { background:#1c2232; }
.name { font-weight:700; }
.badge { display:inline-block; padding:1px 7px; border-radius:999px; font-size:11px;
  margin-right:4px; }
.el-fire { background:#e85d4c33; color:var(--fire); }
.el-water { background:#4aa3e833; color:var(--water); }
.el-wood { background:#5cbf6a33; color:var(--wood); }
.el-light { background:#e8c84a33; color:var(--light); }
.el-dark { background:#9b7ae833; color:var(--dark); }
.el-none { background:#ffffff14; color:var(--muted); }
.rar { color:#d4b15f; font-size:12px; }
.av { width:64px; height:64px; object-fit:cover; border-radius:8px;
  background:#0b0e16; display:block; }
.av.sm { width:88px; height:88px; }
.ph { width:64px; height:64px; border-radius:8px; background:#2a3144;
  display:flex; align-items:center; justify-content:center; color:var(--muted); font-size:12px; }
.stats { display:grid; grid-template-columns:repeat(3,auto); gap:2px 10px; font-size:12px; }
.stats.na { opacity:.4; }
.stats i { color:var(--muted); font-style:normal; margin-right:4px; }
.sk { font-size:12px; color:#d5dbe8; max-width:280px; max-height:4.2em;
  overflow:hidden; white-space:pre-wrap; }
.sk.ign { color:#ff8b8b; }
.empty { color:#555d72; }
.gallery { display:grid; grid-template-columns:repeat(auto-fill,minmax(148px,1fr)); gap:10px; }
.card {
  background:var(--panel); border:1px solid var(--line); border-radius:12px;
  padding:10px; text-align:center;
}
.card img { width:100%; aspect-ratio:1; object-fit:cover; border-radius:10px; background:#0b0e16; }
.card .nm { margin-top:6px; font-weight:700; font-size:13px; }
.hidden { display:none !important; }
footer { color:var(--muted); font-size:12px; padding:8px 16px 24px; }
"""

JS = r"""
function ph(){ const d=document.createElement('div'); d.className='ph'; d.textContent='无'; return d; }
const $ = (s, r=document) => r.querySelector(s);
const $$ = (s, r=document) => [...r.querySelectorAll(s)];
function showPanel(id){
  $$('.panel').forEach(p=>p.classList.toggle('on', p.id===id));
  $$('.tabs button').forEach(b=>b.classList.toggle('on', b.dataset.panel===id));
  filterAll();
}
function setView(v){
  $$('.views button').forEach(b=>b.classList.toggle('on', b.dataset.view===v));
  $('#char-table-wrap').classList.toggle('hidden', v!=='table');
  $('#char-gallery').classList.toggle('hidden', v!=='gallery');
}
function hay(el){ return (el.dataset.hay||''); }
function match(el, q, f){
  if(q && !hay(el).includes(q)) return false;
  if(f.el && el.dataset.element !== f.el) return false;
  if(f.role && el.dataset.role !== f.role) return false;
  if(f.rar && el.dataset.rarity !== f.rar) return false;
  return true;
}
function filterAll(){
  const q = ($('#q').value||'').trim().toLowerCase();
  const f = { el:$('#fel').value, role:$('#frole').value, rar:$('#frar').value };
  let n=0, t=0;
  $$('#char-table tbody tr').forEach(tr=>{
    t++; const ok=match(tr,q,f); tr.classList.toggle('hidden', !ok); if(ok) n++;
  });
  $$('#char-gallery .card').forEach(c=>{
    const ok=match(c,q,f); c.classList.toggle('hidden', !ok);
  });
  const cartaQ = q;
  let cn=0;
  $$('#carta-table tbody tr').forEach(tr=>{
    const ok=!cartaQ || hay(tr).includes(cartaQ);
    tr.classList.toggle('hidden', !ok); if(ok) cn++;
  });
  let pn=0;
  $$('#puppet-table tbody tr').forEach(tr=>{
    const ok=!cartaQ || hay(tr).includes(cartaQ);
    tr.classList.toggle('hidden', !ok); if(ok) pn++;
  });
  $('#meta').textContent = '天子 '+n+'/'+t+'　魂卡显示 '+cn+'　人偶显示 '+pn;
}
document.addEventListener('input', e=>{
  if(['q','fel','frole','frar'].includes(e.target.id)) filterAll();
});
"""


def build_html(chars, cartas, puppets, buffs, stats_note: str) -> str:
    el_opts = "".join('<option value="%s">%s</option>' % (x, x) for x in ("火", "水", "木", "光", "暗"))
    role_opts = "".join(
        '<option value="%s">%s</option>' % (x, x)
        for x in ("攻击型", "防御型", "干扰型", "治疗型", "辅助型")
    )
    rar_opts = "".join('<option value="%s">%s</option>' % (x, x) for x in ("5星", "4星"))

    char_rows = []
    cards = []
    for r in chars:
        name = r.get("name") or ""
        el = r.get("element") or ""
        role = r.get("role") or ""
        rar = r.get("rarity") or ""
        hay = " ".join([
            name, el, role, rar,
            r.get("auto_text") or "", r.get("tap_text") or "", r.get("slide_text") or "",
            r.get("drive_text") or "", r.get("leader_text") or "",
        ]).lower()
        wiki = r.get("wiki_url") or ""
        ds = 'data-element="%s" data-role="%s" data-rarity="%s" data-hay="%s"' % (
            e(el), e(role), e(rar), e(hay)
        )
        char_rows.append(
            "<tr %s><td>%s</td><td><div class='name'>%s</div>"
            "<div><span class='badge %s'>%s</span>"
            "<span class='badge el-none'>%s</span>"
            "<span class='rar'>%s</span></div>"
            "%s</td><td><div class='lbl'>初始</div>%s"
            "<div class='lbl'>满破</div>%s</td>"
            "<td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
            % (
                ds,
                avatar_html(r),
                e(name),
                el_class(el), e(el or "属性?"),
                e(role or "职业?"),
                e(rar),
                ('<div><a href="%s" target="_blank">wiki</a></div>' % e(wiki)) if wiki else "",
                stat_block(r, "init"),
                stat_block(r, "max"),
                skill_cell(r.get("auto_text")),
                skill_cell(r.get("tap_text"), r.get("tap_ignited")),
                skill_cell(r.get("slide_text"), r.get("slide_ignited")),
                skill_cell(r.get("drive_text"), r.get("drive_ignited")),
                skill_cell(r.get("leader_text")),
            )
        )
        img = r.get("avatar_file") or r.get("avatar_url") or ""
        img_html = (
            '<img src="%s" alt="%s" loading="lazy">' % (e(img), e(name))
            if img else '<div class="ph" style="width:100%;aspect-ratio:1">无头像</div>'
        )
        cards.append(
            '<div class="card" %s>%s<div class="nm">%s</div>'
            '<div><span class="badge %s">%s</span>'
            '<span class="badge el-none">%s</span></div></div>'
            % (ds, img_html, e(name), el_class(el), e(el or "?"), e(role or rar))
        )

    carta_rows = []
    for r in cartas:
        hay = " ".join([r.get("name") or "", r.get("special") or "", r.get("element_gate") or ""]).lower()
        carta_rows.append(
            "<tr data-hay='%s'><td>%s</td><td>%s</td><td>%s</td><td>%s</td>"
            "<td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
            % (
                e(hay),
                avatar_html(r),
                e(r.get("name")),
                e(r.get("rarity")),
                e(r.get("stat_pair")),
                e(r.get("element_gate")),
                e(r.get("role_gate")),
                e(r.get("mode_gate")),
                e(r.get("collab")),
                e(r.get("special")),
            )
        )

    puppet_rows = []
    for r in puppets:
        hay = " ".join([r.get("name") or "", r.get("tap") or "", r.get("slide") or "", r.get("element") or ""]).lower()
        puppet_rows.append(
            "<tr data-hay='%s'><td>%s</td><td>%s</td><td>%s</td><td>%s</td>"
            "<td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
            % (
                e(hay),
                avatar_html(r),
                e(r.get("name")),
                e(r.get("element")),
                e(r.get("rarity")),
                e(r.get("tap")),
                e(r.get("slide")),
                e(r.get("drive")),
                e(r.get("leader")),
            )
        )

    buff_rows = []
    for r in buffs:
        buff_rows.append(
            "<tr><td>%s</td><td>%s</td><td>%s</td></tr>"
            % (e(r.get("kind")), e(r.get("name")), e(r.get("effect")))
        )

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>天命之子 · 数据汇总表</title>
<style>{CSS}
.lbl {{ color:var(--muted); font-size:11px; margin-top:4px; }}
</style>
</head>
<body>
<header>
  <h1>天命之子数据汇总</h1>
  <div class="sub">GameKee 词条整理 · 天子 {len(chars)} · 魂之歌牌 {len(cartas)} · 人偶 {len(puppets)} · Buff {len(buffs)}
  <br>{e(stats_note)}</div>
  <div class="tabs">
    <button class="on" data-panel="p-child" onclick="showPanel('p-child')">天子</button>
    <button data-panel="p-carta" onclick="showPanel('p-carta')">魂之歌牌</button>
    <button data-panel="p-puppet" onclick="showPanel('p-puppet')">人偶</button>
    <button data-panel="p-buff" onclick="showPanel('p-buff')">Buff</button>
  </div>
  <div class="filters">
    <input id="q" placeholder="搜索名字 / 技能 / 效果">
    <select id="fel"><option value="">全部属性</option>{el_opts}</select>
    <select id="frole"><option value="">全部职业</option>{role_opts}</select>
    <select id="frar"><option value="">全部稀有度</option>{rar_opts}</select>
    <span class="views">
      <button class="on" data-view="table" onclick="setView('table')">表格</button>
      <button data-view="gallery" onclick="setView('gallery')">头像图鉴</button>
    </span>
    <span class="meta" id="meta"></span>
  </div>
</header>
<main>
<section id="p-child" class="panel on">
  <div id="char-table-wrap" class="table-wrap">
    <table id="char-table">
      <thead><tr>
        <th>头像</th><th>天子</th><th>面板</th>
        <th>普攻</th><th>重击 TS</th><th>滑动 SS</th><th>大招 DS</th><th>队长技</th>
      </tr></thead>
      <tbody>
        {''.join(char_rows)}
      </tbody>
    </table>
  </div>
  <div id="char-gallery" class="gallery hidden">
    {''.join(cards)}
  </div>
</section>
<section id="p-carta" class="panel">
  <div class="table-wrap">
    <table id="carta-table">
      <thead><tr>
        <th>图</th><th>名称</th><th>稀有度</th><th>配对</th><th>属性</th>
        <th>职业</th><th>场合</th><th>联动</th><th>特效</th>
      </tr></thead>
      <tbody>{''.join(carta_rows)}</tbody>
    </table>
  </div>
</section>
<section id="p-puppet" class="panel">
  <div class="table-wrap">
    <table id="puppet-table">
      <thead><tr>
        <th>图</th><th>名称</th><th>属性</th><th>稀有度</th>
        <th>重击</th><th>滑动</th><th>大招</th><th>队长</th>
      </tr></thead>
      <tbody>{''.join(puppet_rows)}</tbody>
    </table>
  </div>
</section>
<section id="p-buff" class="panel">
  <div class="table-wrap">
    <table>
      <thead><tr><th>类型</th><th>名称</th><th>效果</th></tr></thead>
      <tbody>{''.join(buff_rows)}</tbody>
    </table>
  </div>
</section>
</main>
<footer>头像来自 GameKee 词条本地缓存，仅供非商业复刻资料整理。用浏览器打开本 HTML 即可看图；CSV 可用 Excel / WPS 打开。</footer>
<script>{JS}
filterAll();
</script>
</body>
</html>
"""


def try_xlsx(chars, cartas, puppets, buffs, avatars_dir, dest):
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from PIL import Image as PILImage
    except ImportError:
        return False, "no-openpyxl"

    thumb_dir = os.path.join(avatars_dir, "_xlsx_thumbs")
    os.makedirs(thumb_dir, exist_ok=True)

    def thumb(rec):
        src = rec.get("avatar_file") or ""
        if not src:
            return None
        src_path = os.path.join(os.path.dirname(avatars_dir), src) if not os.path.isabs(src) else src
        # avatar_file is 'avatars/123.png' relative to workbook folder
        alt = os.path.join(avatars_dir, os.path.basename(src))
        if os.path.isfile(alt):
            src_path = alt
        if not os.path.isfile(src_path):
            return None
        out = os.path.join(thumb_dir, os.path.splitext(os.path.basename(src_path))[0] + ".png")
        if not os.path.isfile(out):
            im = PILImage.open(src_path).convert("RGBA")
            im.thumbnail((64, 64))
            canvas = PILImage.new("RGBA", (64, 64), (16, 19, 26, 255))
            x = (64 - im.size[0]) // 2
            y = (64 - im.size[1]) // 2
            canvas.paste(im, (x, y), im)
            canvas.convert("RGB").save(out, "PNG")
        return out

    wb = Workbook()

    def style_header(ws, n):
        fill = PatternFill("solid", fgColor="1E2433")
        font = Font(color="E8ECF4", bold=True)
        for col in range(1, n + 1):
            cell = ws.cell(1, col)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 22

    ws = wb.active
    ws.title = "天子汇总"
    headers = [
        "头像", "content_id", "名称", "稀有度", "属性", "职业",
        "初始CP", "初始HP", "初始攻", "初始防", "初始敏", "初始暴",
        "满破CP", "满破HP", "满破攻", "满破防", "满破敏", "满破暴",
        "普攻", "重击TS", "TS点火", "滑动SS", "SS点火", "大招DS", "DS点火", "队长技",
        "头像文件", "头像URL", "wiki",
    ]
    ws.append(headers)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["C"].width = 18
    for col in range(19, 27):
        ws.column_dimensions[get_column_letter(col)].width = 28
    for i, r in enumerate(chars, start=2):
        ws.append([
            "",
            int(r["content_id"]) if str(r.get("content_id") or "").isdigit() else r.get("content_id"),
            r.get("name"), r.get("rarity"), r.get("element"), r.get("role"),
            r.get("cp_init"), r.get("hp_init"), r.get("atk_init"), r.get("def_init"),
            r.get("agl_init"), r.get("crt_init"),
            r.get("cp_max"), r.get("hp_max"), r.get("atk_max"), r.get("def_max"),
            r.get("agl_max"), r.get("crt_max"),
            r.get("auto_text"), r.get("tap_text"), r.get("tap_ignited"),
            r.get("slide_text"), r.get("slide_ignited"),
            r.get("drive_text"), r.get("drive_ignited"), r.get("leader_text"),
            r.get("avatar_file"), r.get("avatar_url"), r.get("wiki_url"),
        ])
        ws.row_dimensions[i].height = 52
        tp = thumb(r)
        if tp:
            img = XLImage(tp)
            img.width = 48
            img.height = 48
            ws.add_image(img, "A%d" % i)
        for col in range(1, len(headers) + 1):
            ws.cell(i, col).alignment = Alignment(wrap_text=True, vertical="center")
    style_header(ws, len(headers))

    ws2 = wb.create_sheet("魂之歌牌")
    h2 = ["头像", "content_id", "名称", "稀有度", "配对", "属性", "职业", "场合", "联动", "特效", "头像文件", "头像URL"]
    ws2.append(h2)
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["J"].width = 50
    for i, r in enumerate(cartas, start=2):
        ws2.append([
            "", r.get("content_id"), r.get("name"), r.get("rarity"), r.get("stat_pair"),
            r.get("element_gate"), r.get("role_gate"), r.get("mode_gate"), r.get("collab"),
            r.get("special"), r.get("avatar_file"), r.get("avatar_url"),
        ])
        ws2.row_dimensions[i].height = 52
        tp = thumb(r)
        if tp:
            img = XLImage(tp)
            img.width = 48
            img.height = 48
            ws2.add_image(img, "A%d" % i)
    style_header(ws2, len(h2))

    ws3 = wb.create_sheet("人偶")
    h3 = ["头像", "content_id", "名称", "属性", "稀有度", "重击", "滑动", "大招", "队长", "头像文件", "头像URL"]
    ws3.append(h3)
    ws3.column_dimensions["A"].width = 12
    ws3.column_dimensions["F"].width = 36
    ws3.column_dimensions["G"].width = 36
    for i, r in enumerate(puppets, start=2):
        ws3.append([
            "", r.get("content_id"), r.get("name"), r.get("element"), r.get("rarity"),
            r.get("tap"), r.get("slide"), r.get("drive"), r.get("leader"),
            r.get("avatar_file"), r.get("avatar_url"),
        ])
        ws3.row_dimensions[i].height = 52
        tp = thumb(r)
        if tp:
            img = XLImage(tp)
            img.width = 48
            img.height = 48
            ws3.add_image(img, "A%d" % i)
    style_header(ws3, len(h3))

    ws4 = wb.create_sheet("Buff")
    ws4.append(["类型", "名称", "效果"])
    for r in buffs:
        ws4.append([r.get("kind"), r.get("name"), r.get("effect")])
    style_header(ws4, 3)
    ws4.column_dimensions["C"].width = 50

    wb.save(dest)
    return True, dest


def copy_tree_files(files, dest_dir):
    os.makedirs(dest_dir, exist_ok=True)
    for src in files:
        shutil.copy2(src, os.path.join(dest_dir, os.path.basename(src)))


def main() -> int:
    char_csv = os.path.join(TABLES, "characters.csv")
    if not os.path.isfile(char_csv):
        print("missing", char_csv)
        return 1

    chars = read_csv(char_csv)
    cartas = read_csv(os.path.join(TABLES, "soul_cartas.csv"))
    puppets = read_csv(os.path.join(TABLES, "puppets.csv"))
    buffs = read_csv(os.path.join(TABLES, "buffs.csv"))
    print("loaded chars %d cartas %d puppets %d buffs %d" % (
        len(chars), len(cartas), len(puppets), len(buffs)))

    print("scan related_strategies icons…", flush=True)
    rel_id, rel_name = collect_related_icons()
    print("  related icons by id %d by name %d" % (len(rel_id), len(rel_name)))

    src_counter = Counter()
    for rec in chars:
        cid = int(rec["content_id"])
        url, src = pick_avatar(cid, rec.get("name") or "", rel_id, rel_name)
        rec["avatar_url"] = url
        rec["wiki_url"] = "https://www.gamekee.com/dc/%s.html" % cid
        rec["_src"] = src
        src_counter[src] += 1
    print("character avatar sources:", dict(src_counter))

    for group in (cartas, puppets):
        for rec in group:
            cid = rec.get("content_id") or ""
            if not cid:
                rec["avatar_url"] = ""
                continue
            cid_i = int(cid)
            url, src = pick_avatar(cid_i, rec.get("name") or "", rel_id, rel_name)
            rec["avatar_url"] = url
            rec["_src"] = src

    avatars_user = os.path.join(USER_OUT, "avatars")
    avatars_ref = os.path.join(TABLES, "avatars")
    os.makedirs(avatars_user, exist_ok=True)
    os.makedirs(avatars_ref, exist_ok=True)

    print("download character portraits…", flush=True)
    st1 = download_all(chars, avatars_user)
    print("  chars", dict(st1))
    print("download carta portraits…", flush=True)
    st2 = download_all(cartas, avatars_user)
    print("  cartas", dict(st2))
    print("download puppet portraits…", flush=True)
    st3 = download_all(puppets, avatars_user)
    print("  puppets", dict(st3))

    # mirror avatars into reference tables
    if os.path.isdir(avatars_user):
        for fn in os.listdir(avatars_user):
            src = os.path.join(avatars_user, fn)
            if os.path.isfile(src):
                shutil.copy2(src, os.path.join(avatars_ref, fn))

    have_av = sum(1 for r in chars if r.get("avatar_file"))
    have_url = sum(1 for r in chars if r.get("avatar_url"))
    note = "天子头像已下载 %d/%d，找到 URL %d/%d；满破面板空着是词条本身没填。" % (
        have_av, len(chars), have_url, len(chars)
    )

    html_path = os.path.join(USER_OUT, "汇总表.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(build_html(chars, cartas, puppets, buffs, note))
    shutil.copy2(html_path, os.path.join(TABLES, "汇总表.html"))

    write_csv(os.path.join(USER_OUT, "汇总表.csv"), chars, CHAR_FIELDS)
    write_csv(os.path.join(TABLES, "汇总表.csv"), chars, CHAR_FIELDS)

    carta_fields = ["content_id", "avatar_file", "avatar_url"] + [
        k for k in (cartas[0].keys() if cartas else []) if k not in ("avatar_file", "avatar_url", "_src")
    ]
    puppet_fields = ["content_id", "avatar_file", "avatar_url"] + [
        k for k in (puppets[0].keys() if puppets else []) if k not in ("avatar_file", "avatar_url", "_src")
    ]
    # de-dup while keeping order
    def uniq(seq):
        seen = set()
        out = []
        for x in seq:
            if x in seen:
                continue
            seen.add(x)
            out.append(x)
        return out

    write_csv(os.path.join(USER_OUT, "魂之歌牌.csv"), cartas, uniq(carta_fields))
    write_csv(os.path.join(USER_OUT, "人偶.csv"), puppets, uniq(puppet_fields))

    xlsx_path = os.path.join(USER_OUT, "汇总表.xlsx")
    ok, info = try_xlsx(chars, cartas, puppets, buffs, avatars_user, xlsx_path)
    print("xlsx", ok, info)
    if ok:
        shutil.copy2(xlsx_path, os.path.join(TABLES, "汇总表.xlsx"))

    missing = [r["name"] for r in chars if not r.get("avatar_url")]
    print("missing avatar URL:", len(missing), missing[:20])
    nofile = [r["name"] for r in chars if r.get("avatar_url") and not r.get("avatar_file")]
    print("download failed:", len(nofile), nofile[:20])
    print("HTML", html_path)
    print("CSV", os.path.join(USER_OUT, "汇总表.csv"))
    print("avatars", avatars_user, "count", len([x for x in os.listdir(avatars_user) if os.path.isfile(os.path.join(avatars_user, x))]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
