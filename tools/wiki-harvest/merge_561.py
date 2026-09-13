# -*- coding: utf-8 -*-
"""Merge Memorial 561 OCR + GameKee wiki + grid portraits into 汇总表."""
from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
from collections import defaultdict
from difflib import SequenceMatcher

from zhconv import convert as zh_convert

ROOT = r"F:\天命之子"
USER = os.path.join(ROOT, "天命之子数据")
EXT = os.path.join(USER, "_extract")
WIKI_CHAR = os.path.join(ROOT, "docs", "reference", "gamekee", "tables", "characters.csv")
TABLES = os.path.join(ROOT, "docs", "reference", "gamekee", "tables")

CJK_SPACE = re.compile(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])")


def simp(s: str) -> str:
    s = CJK_SPACE.sub("", s or "")
    return zh_convert(re.sub(r"\s+", " ", s), "zh-cn").strip()


def load_jsonl(*names):
    rows = []
    for n in names:
        p = os.path.join(EXT, n)
        if not os.path.isfile(p):
            print("missing", p)
            continue
        with open(p, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or not line.startswith("{"):
                    continue
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    return rows


def core_key(name: str) -> str:
    n = simp(name)
    n = n.replace("的", "")
    n = re.sub(r"[♥♡★☆\s·・]", "", n)
    return n


def load_wiki():
    if not os.path.isfile(WIKI_CHAR):
        return []
    with open(WIKI_CHAR, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def looks_like_quote(name: str) -> bool:
    n = name or ""
    if len(n) >= 12:
        return True
    if any(ch in n for ch in "。？！?!.…"):
        return True
    return False


def garbage_name(name: str) -> bool:
    n = (name or "").strip()
    if len(n) < 2:
        return True
    if looks_like_quote(n):
        return True
    if len(set(n)) <= 2 and len(n) >= 4:
        return True
    if re.fullmatch(r"[一二三四五六七八九十0-9]+", n) and len(n) >= 3:
        return True
    return False


def rar_key(s: str) -> str:
    t = str(s or "")
    for ch in "54321":
        if t.startswith(ch):
            return ch
    return ""


def pick_name(rec: dict) -> str:
    name = simp(rec.get("name") or "")
    ocr = rec.get("name_ocr") or ""
    parts = [simp(p) for p in ocr.split("|")]
    cv = ""
    short = []
    for p in parts:
        if p.upper().startswith("CV") or p.startswith("cv"):
            cv = p
            continue
        if not p or "属性" in p or "战斗" in p or "HP" in p:
            continue
        if looks_like_quote(p):
            continue
        if 2 <= len(p) <= 12 and not garbage_name(p):
            short.append(p)
    if short:
        name = max(short, key=len)
    elif looks_like_quote(name):
        name = ""
    rec["_cv"] = rec.get("cv") or cv
    return name


def wiki_score(mem_name: str, wiki_name: str) -> int:
    cn = core_key(mem_name)
    wc = core_key(wiki_name)
    if not cn or not wc or len(cn) < 2:
        return 0
    if wc == cn:
        return 200
    if cn.endswith(wc) or wc.endswith(cn):
        if min(len(cn), len(wc)) >= 2:
            return 120 + min(len(cn), len(wc))
    if len(cn) >= 3 and (cn in wc or wc in cn):
        return 90 + min(len(cn), len(wc))
    for n in range(min(len(cn), len(wc), 6), 2, -1):
        if cn[-n:] == wc[-n:]:
            return 50 + n * 8
    return 0


def assign_wiki(mem_rows, wiki_rows):
    """Highest-score unique assignment; skip quotes and weak pairs."""
    pairs = []
    for i, rec in enumerate(mem_rows):
        names = [rec.get("name") or ""]
        for p in (rec.get("name_ocr") or "").split("|"):
            names.append(simp(p))
        rk = rar_key(rec.get("rarity") or rec.get("rarity_key") or "")
        best_s = 0
        best_w = None
        for name in names:
            if garbage_name(name):
                continue
            for w in wiki_rows:
                if rk and rar_key(w.get("rarity") or "") not in ("", rk):
                    continue
                s = wiki_score(name, w.get("name") or "")
                if s > best_s:
                    best_s, best_w = s, w
        if best_s >= 70 and best_w is not None:
            pairs.append((best_s, i, id(best_w), best_w))
    pairs.sort(key=lambda x: -x[0])
    used_m, used_w = set(), set()
    mapping = {}
    for s, i, wid, w in pairs:
        if i in used_m or wid in used_w:
            continue
        used_m.add(i)
        used_w.add(wid)
        mapping[i] = w
    return mapping


def merge():
    mem = load_jsonl("5a.jsonl", "5b.jsonl", "4.jsonl", "3all.jsonl", "2.jsonl", "1.jsonl")
    for ov in load_jsonl("3sk_a.jsonl", "3sk_b.jsonl", "2sk.jsonl", "1sk.jsonl", "5sk.jsonl", "4sk.jsonl"):
        key = (str(ov.get("rarity_key") or ""), ov.get("idx"))
        for r in mem:
            if (str(r.get("rarity_key") or ""), r.get("idx")) == key:
                for f in ("auto_text", "tap_text", "slide_text", "drive_text", "leader_text", "skills_ocr"):
                    if len(ov.get(f) or "") > len(r.get(f) or ""):
                        r[f] = ov.get(f) or ""
                break
    print("memorial rows", len(mem))
    wiki = load_wiki()
    grid = load_jsonl("grid_avatars.jsonl")

    def richness(r: dict) -> int:
        return (
            len(r.get("tap_text") or "")
            + len(r.get("slide_text") or "")
            + len(r.get("drive_text") or "")
            + len(r.get("skills_ocr") or "")
            + len(r.get("hp") or "")
            + len(r.get("name") or "")
        )

    by_best = {}
    for r in mem:
        raw = simp(r.get("name") or "")
        r["name"] = pick_name(r) if garbage_name(raw) else raw
        key = str(r.get("rarity_key") or "")
        try:
            ix = int(r.get("idx"))
        except (TypeError, ValueError):
            continue
        old = by_best.get((key, ix))
        if old is None or richness(r) > richness(old):
            by_best[(key, ix)] = r

    ordered = []
    for key in ("5", "4", "3", "2", "1"):
        idxs = sorted(i for (k, i) in by_best if k == key)
        for ix in idxs:
            ordered.append(by_best[(key, ix)])
    wiki_map = assign_wiki(ordered, wiki)
    used_w = {id(w) for w in wiki_map.values()}
    for i, rec in enumerate(ordered):
        if i in wiki_map:
            continue
        if str(rec.get("rarity_key") or rec.get("rarity") or "") not in ("5", "4", "5星", "4星"):
            continue
        name = rec.get("name") or ""
        if garbage_name(name):
            continue
        best_s, best_w = 0, None
        for w in wiki:
            if id(w) in used_w:
                continue
            s = wiki_score(name, w.get("name") or "")
            if s > best_s:
                best_s, best_w = s, w
        if best_s < 66:
            for w in wiki:
                if id(w) in used_w:
                    continue
                ratio = SequenceMatcher(None, core_key(name), core_key(w.get("name") or "")).ratio()
                s = int(ratio * 100)
                if s > best_s:
                    best_s, best_w = s, w
        if best_s >= 62 and best_w is not None:
            wiki_map[i] = best_w
            used_w.add(id(best_w))

    out = []
    for i, r in enumerate(ordered):
        ix = r.get("idx")
        name = simp(r.get("name") or "")
        rec = {
                "id": r.get("id") or "",
                "idx": ix,
                "name": name,
                "rarity": r.get("rarity") or "",
                "element": r.get("element") or "",
                "role": r.get("role") or "",
                "cp": r.get("cp") or "",
                "hp": r.get("hp") or "",
                "atk": r.get("atk") or "",
                "def_": r.get("def") or "",
                "agl": r.get("agl") or "",
                "crt": r.get("crt") or "",
                "auto_text": r.get("auto_text") or "",
                "tap_text": r.get("tap_text") or "",
                "slide_text": r.get("slide_text") or "",
                "drive_text": r.get("drive_text") or "",
                "leader_text": r.get("leader_text") or "",
                "cv": r.get("cv") or r.get("_cv") or "",
                "quote": r.get("quote") or "",
                "has_stats_panel": r.get("has_stats_panel"),
                "wiki_id": "",
                "wiki_name": "",
                "avatar_file": "",
                "source": "memorial",
            }
        blob = " ".join([r.get("stats_ocr") or "", r.get("element") or "", r.get("role") or "", r.get("skills_ocr") or ""])
        if not rec["element"]:
            for tok, val in (("闇", "暗"), ("暗", "暗"), ("火", "火"), ("水", "水"), ("木", "木"), ("光", "光")):
                if tok in blob:
                    rec["element"] = val
                    break
        if not rec["role"]:
            for tok, val in (
                ("攻擊", "攻击型"), ("攻击", "攻击型"),
                ("防禦", "防御型"), ("防御", "防御型"), ("防型", "防御型"),
                ("妨害", "干扰型"), ("干擾", "干扰型"), ("干扰", "干扰型"), ("干型", "干扰型"),
                ("治癒", "治疗型"), ("治療", "治疗型"), ("治疗", "治疗型"), ("治瘾", "治疗型"),
                ("輔助", "辅助型"), ("辅助", "辅助型"),
            ):
                if tok in blob:
                    rec["role"] = val
                    break
        w = wiki_map.get(i)
        if w:
            rec["wiki_id"] = w.get("content_id") or ""
            rec["wiki_name"] = w.get("name") or ""
            rec["source"] = "memorial+wiki"
            if rec["wiki_name"] and (not rec["name"] or len(rec["wiki_name"]) >= 2):
                rec["name"] = rec["wiki_name"]
            if not rec["element"]:
                rec["element"] = w.get("element") or ""
            if not rec["role"]:
                rec["role"] = w.get("role") or ""
            for a, b in (
                ("auto_text", "auto_text"),
                ("tap_text", "tap_text"),
                ("slide_text", "slide_text"),
                ("drive_text", "drive_text"),
                ("leader_text", "leader_text"),
            ):
                wt = w.get(b) or ""
                if len(wt) > len(rec[a]):
                    rec[a] = wt
            rec["tap_ignited"] = w.get("tap_ignited") or ""
            rec["slide_ignited"] = w.get("slide_ignited") or ""
            rec["drive_ignited"] = w.get("drive_ignited") or ""
            rec["hp_init"] = w.get("hp_init") or rec["hp"]
            rec["atk_init"] = w.get("atk_init") or rec["atk"]
            rec["def_init"] = w.get("def_init") or rec["def_"]
            rec["agl_init"] = w.get("agl_init") or rec["agl"]
            rec["crt_init"] = w.get("crt_init") or rec["crt"]
            rec["hp_max"] = w.get("hp_max") or ""
            rec["atk_max"] = w.get("atk_max") or ""
            rec["def_max"] = w.get("def_max") or ""
            rec["agl_max"] = w.get("agl_max") or ""
            rec["crt_max"] = w.get("crt_max") or ""
            av = os.path.join(USER, "avatars", "%s.png" % w.get("content_id"))
            jpg = os.path.join(USER, "avatars", "%s.jpg" % w.get("content_id"))
            if os.path.isfile(av):
                rec["avatar_file"] = "avatars/%s.png" % w.get("content_id")
            elif os.path.isfile(jpg):
                rec["avatar_file"] = "avatars/%s.jpg" % w.get("content_id")
        else:
            rec["tap_ignited"] = rec["slide_ignited"] = rec["drive_ignited"] = ""
            rec["hp_init"] = rec["hp"]
            rec["atk_init"] = rec["atk"]
            rec["def_init"] = rec["def_"]
            rec["agl_init"] = rec["agl"]
            rec["crt_init"] = rec["crt"]
            rec["hp_max"] = rec["atk_max"] = rec["def_max"] = rec["agl_max"] = rec["crt_max"] = ""
        out.append(rec)

    # Catalog card icons: match by OCR'd tile name, never by swipe index for 5/4.
    named = load_jsonl("named_icons.jsonl")
    used_icon = set()

    def icon_score(char_name: str, grid_name: str) -> int:
        s = wiki_score(char_name, grid_name)
        if s:
            return s
        gn = core_key(grid_name)
        cn = core_key(char_name)
        if not gn or not cn:
            return 0
        return int(SequenceMatcher(None, cn, gn).ratio() * 100)

    for rec in out:
        rar = str(rec.get("rarity") or "")[:1]
        keep_wiki = rec["avatar_file"].startswith("avatars/") and rec.get("wiki_id")
        # 5/4 keep wiki heads; still try named catalog card if current file is idx-based or screenshot
        current = rec["avatar_file"] or ""
        idx_based = "/R%s_" % rar in current or current.startswith("icons/R")
        if keep_wiki and not idx_based:
            continue
        best, best_s = None, 0
        for g in named:
            if g.get("rarity_key") != rar:
                continue
            if g.get("file") in used_icon:
                continue
            s = icon_score(rec.get("name") or "", g.get("name_grid") or "")
            if s > best_s:
                best_s, best = s, g
        if best and best_s >= 55:
            rec["avatar_file"] = best["file"]
            used_icon.add(best["file"])
            continue
        if rar in ("2", "1"):
            ix = rec.get("idx")
            icon = os.path.join(USER, "icons", "R%s_%03d.png" % (rar, ix if ix is not None else -1))
            if os.path.isfile(icon):
                rec["avatar_file"] = "icons/R%s_%03d.png" % (rar, ix)
        elif not rec["avatar_file"] or idx_based or rec["avatar_file"].startswith("avatars_561/"):
            rec["avatar_file"] = rec["avatar_file"] if keep_wiki else rec["avatar_file"]
            if (not rec["avatar_file"]) or idx_based or rec["avatar_file"].startswith("avatars_561/"):
                if keep_wiki:
                    pass
                else:
                    rec["avatar_file"] = rec.get("avatar_file") or ""

    return out, wiki, grid


def e(s):
    return html.escape("" if s is None else str(s), quote=True)


def write_csv(path, rows):
    fields = [
        "id", "idx", "name", "wiki_name", "wiki_id", "rarity", "element", "role",
        "avatar_file", "cp", "hp", "atk", "def_", "agl", "crt",
        "hp_init", "atk_init", "def_init", "agl_init", "crt_init",
        "hp_max", "atk_max", "def_max", "agl_max", "crt_max",
        "auto_text", "tap_text", "tap_ignited", "slide_text", "slide_ignited",
        "drive_text", "drive_ignited", "leader_text", "cv", "source",
    ]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_html(path, rows):
    from gallery_html import apply_face_icons, write_html as write_gallery

    apply_face_icons(rows)
    write_gallery(path, rows)


def main():
    rows, wiki, grid = merge()
    print("merged", len(rows), "wiki matched", sum(1 for r in rows if r.get("wiki_id")),
          "avatars", sum(1 for r in rows if r.get("avatar_file")))
    byr = defaultdict(int)
    for r in rows:
        byr[r.get("rarity")] += 1
    print("by rarity", dict(byr))
    csv_path = os.path.join(USER, "汇总表.csv")
    html_path = os.path.join(USER, "汇总表.html")
    write_csv(csv_path, rows)
    write_html(html_path, rows)
    shutil.copy2(csv_path, os.path.join(TABLES, "汇总表.csv"))
    shutil.copy2(html_path, os.path.join(TABLES, "汇总表.html"))
    print("wrote", html_path, csv_path)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "天子561"
        fields = [
            "id", "idx", "name", "wiki_name", "wiki_id", "rarity", "element", "role",
            "avatar_file", "cp", "hp", "atk", "def_", "agl", "crt",
            "hp_init", "atk_init", "def_init", "agl_init", "crt_init",
            "hp_max", "atk_max", "def_max", "agl_max", "crt_max",
            "auto_text", "tap_text", "tap_ignited", "slide_text", "slide_ignited",
            "drive_text", "drive_ignited", "leader_text", "cv", "source",
        ]
        ws.append(fields)
        fill = PatternFill("solid", fgColor="1E2433")
        for col in range(1, len(fields) + 1):
            ws.cell(1, col).fill = fill
            ws.cell(1, col).font = Font(color="E8ECF4", bold=True)
        for r in rows:
            ws.append([r.get(k, "") for k in fields])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        # extra sheets from previous tables
        for title, src in (("魂之歌牌", os.path.join(USER, "魂之歌牌.csv")),
                           ("人偶", os.path.join(USER, "人偶.csv")),
                           ("Buff", os.path.join(USER, "buffs.csv"))):
            if not os.path.isfile(src):
                continue
            sh = wb.create_sheet(title)
            with open(src, encoding="utf-8-sig") as f:
                rd = csv.reader(f)
                for i, row in enumerate(rd, 1):
                    sh.append(row)
                    if i == 1:
                        for col in range(1, len(row) + 1):
                            sh.cell(1, col).fill = fill
                            sh.cell(1, col).font = Font(color="E8ECF4", bold=True)
            sh.freeze_panes = "A2"
        xlsx_path = os.path.join(USER, "汇总表.xlsx")
        wb.save(xlsx_path)
        shutil.copy2(xlsx_path, os.path.join(TABLES, "汇总表.xlsx"))
        print("wrote", xlsx_path)
    except Exception as e:
        print("xlsx skip", e)


if __name__ == "__main__":
    main()
